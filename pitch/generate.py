"""Generate UsageLens executive pitch deck using python-pptx."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
import os

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Colors
NAVY = RGBColor(0x1E, 0x27, 0x61)
DEEP = RGBColor(0x0D, 0x1B, 0x3E)
ICE = RGBColor(0xCA, 0xDC, 0xFC)
CORAL = RGBColor(0xF9, 0x61, 0x67)
GOLD = RGBColor(0xF9, 0xD4, 0x23)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
OFF_WHITE = RGBColor(0xF5, 0xF7, 0xFA)
GRAY = RGBColor(0x8A, 0x94, 0xA6)
DARK = RGBColor(0x3A, 0x3F, 0x4B)
GREEN = RGBColor(0x2E, 0xCC, 0x71)
TEAL = RGBColor(0x1A, 0xBC, 0x9C)

FONT_T = "Cambria"
FONT_B = "Calibri"

def add_bg(slide, color):
    bg = slide.background
    bg.fill.solid()
    bg.fill.fore_color.rgb = color

def add_text(slide, text, x, y, w, h, size=14, color=WHITE, bold=False, italic=False, font=None, align=PP_ALIGN.LEFT):
    txBox = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.italic = italic
    p.font.name = font or FONT_B
    p.alignment = align
    return txBox

def add_rect(slide, x, y, w, h, color, radius=None):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    if hasattr(shape, 'adjustments') and radius:
        shape.adjustments[0] = radius
    return shape

# =================== SLIDE 1 — TITLE ===================
slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
add_bg(slide, DEEP)

add_text(slide, "UsageLens", 1.0, 1.5, 11, 1.2, size=54, color=WHITE, bold=True, font=FONT_T)
add_text(slide, "Know If Your Loaned Systems Are Being Used", 1.0, 2.9, 9, 0.8, size=26, color=ICE, font=FONT_B)

add_rect(slide, 1.0, 3.9, 2.5, 0.05, CORAL)

add_text(slide, "SSH-Based System Usage Tracking for AMD Field Engineers", 1.0, 4.3, 9, 0.5, size=15, color=GRAY)

# =================== SLIDE 2 — THE PROBLEM ===================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)

add_text(slide, "The Problem", 0.8, 0.4, 6, 0.9, size=40, color=DEEP, bold=True, font=FONT_T)

# Big stat box
add_rect(slide, 0.8, 1.6, 5.5, 2.4, OFF_WHITE, radius=0.05)
add_text(slide, "$100K – $500K+", 1.1, 1.8, 5, 0.8, size=44, color=CORAL, bold=True, font=FONT_T)
add_text(slide, "per GPU system loaned to customers\nfor evaluation (MI300X, MI325X, etc.)", 1.1, 2.7, 5, 0.9, size=15, color=DARK)

# Problem points
problems = [
    ("?", "No visibility into whether customers\nare actually using loaned systems"),
    ("!", "Expensive hardware may sit idle\nfor weeks or months"),
    ("$", "No data to justify continued loans\nor reclaim underutilized systems"),
    ("⏱", "FAEs must ask customers directly —\nsubjective and time-consuming"),
]
for i, (icon, text) in enumerate(problems):
    y = 1.7 + i * 1.2
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(7.0), Inches(y), Inches(0.5), Inches(0.5))
    shape.fill.solid()
    shape.fill.fore_color.rgb = CORAL
    shape.line.fill.background()
    tf = shape.text_frame
    tf.paragraphs[0].text = icon
    tf.paragraphs[0].font.size = Pt(16)
    tf.paragraphs[0].font.color.rgb = WHITE
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER

    add_text(slide, text, 7.7, y, 4.8, 0.55, size=13, color=DARK)

add_text(slide, '"Is the customer even using the system?" — the question every FAE asks', 0.8, 6.4, 11.5, 0.5, size=14, color=GRAY, italic=True)

# =================== SLIDE 3 — THE SOLUTION ===================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, DEEP)

add_text(slide, "The Solution", 0.8, 0.4, 6, 0.9, size=40, color=WHITE, bold=True, font=FONT_T)
add_text(slide, "UsageLens — a lightweight, self-hosted web tool that SSHs into\nremote systems and answers the question with data.", 0.8, 1.3, 11, 0.7, size=16, color=ICE)

# Four capability cards
caps = [
    ("Login Tracking", "Who logged in, when,\nand for how long", "via `last`", TEAL),
    ("GPU Utilization", "Track GPU usage %\nover time", "via rocm-smi", CORAL),
    ("System Load", "CPU load averages\nand uptime trends", "via `uptime`", GOLD),
    ("Zero Install", "No agent on target —\njust SSH access", "SSH only", GREEN),
]
for i, (title, desc, metric, color) in enumerate(caps):
    x = 0.8 + i * 3.05
    add_rect(slide, x, 2.5, 2.75, 3.0, NAVY, radius=0.04)
    add_rect(slide, x, 2.5, 2.75, 0.07, color)
    add_text(slide, title, x + 0.2, 2.85, 2.35, 0.5, size=19, color=WHITE, bold=True, font=FONT_T)
    add_text(slide, desc, x + 0.2, 3.4, 2.35, 0.8, size=13, color=ICE)
    add_text(slide, metric, x + 0.2, 4.6, 2.35, 0.4, size=11, color=color, italic=True)

# =================== SLIDE 4 — KEY FEATURES ===================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)

add_text(slide, "Key Features", 0.8, 0.4, 6, 0.9, size=40, color=DEEP, bold=True, font=FONT_T)

features = [
    ("Collect All Hosts", "One-click collection from every monitored host with per-host status"),
    ("Historical Dashboards", "Pie charts: utilization vs idle — 1 week to 3 months"),
    ("GPU Tracking", "AMD ROCm + NVIDIA — avg and peak utilization over time"),
    ("Jump Host Support", "Reach systems behind firewalls via SSH jump hosts"),
    ("Auto-Polling", "Hands-off monitoring: 1h to 24h intervals, data builds up"),
    ("Snapshot Browser", "Drill into any point in time — users, GPU state, load"),
    ("Login History", "Session durations per user from `last` — real usage data"),
    ("Copy Configs", "Duplicate host setups for quick onboarding of similar systems"),
]
for i, (title, desc) in enumerate(features):
    col = 0 if i < 4 else 1
    row = i % 4
    x = 0.8 + col * 6.2
    y = 1.5 + row * 1.35
    add_rect(slide, x, y, 5.9, 1.1, OFF_WHITE, radius=0.03)
    add_text(slide, title, x + 0.25, y + 0.12, 5.4, 0.35, size=16, color=DEEP, bold=True)
    add_text(slide, desc, x + 0.25, y + 0.55, 5.4, 0.4, size=12, color=GRAY)

# =================== SLIDE 5 — HOW IT WORKS ===================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, OFF_WHITE)

add_text(slide, "How It Works", 0.8, 0.4, 6, 0.9, size=40, color=DEEP, bold=True, font=FONT_T)

# Flow boxes
boxes = [
    ("FAE's Laptop", "Browser UI\nlocalhost:8000", TEAL, 0.8),
    ("Jump Host", "Optional\nfirewall hop", GOLD, 5.1),
    ("Target GPU System", "MI300X / MI325X\nrocm-smi + last", CORAL, 9.4),
]
for label, sub, color, x in boxes:
    add_rect(slide, x, 2.0, 3.0, 2.0, WHITE, radius=0.04)
    add_rect(slide, x, 2.0, 3.0, 0.07, color)
    add_text(slide, label, x + 0.15, 2.3, 2.7, 0.5, size=17, color=DEEP, bold=True, font=FONT_T, align=PP_ALIGN.CENTER)
    add_text(slide, sub, x + 0.15, 2.85, 2.7, 0.7, size=12, color=GRAY, align=PP_ALIGN.CENTER)

# Arrows
for ax in [4.0, 8.3]:
    add_rect(slide, ax, 2.9, 0.9, 0.04, DARK)
    shape = slide.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE, Inches(ax + 0.75), Inches(2.72), Inches(0.3), Inches(0.4))
    shape.fill.solid()
    shape.fill.fore_color.rgb = DARK
    shape.line.fill.background()
    shape.rotation = 90.0

add_text(slide, "SSH", 4.0, 3.15, 0.9, 0.3, size=11, color=GRAY, align=PP_ALIGN.CENTER)
add_text(slide, "SSH", 8.3, 3.15, 0.9, 0.3, size=11, color=GRAY, align=PP_ALIGN.CENTER)

# Tech stack
add_rect(slide, 0.8, 4.8, 11.7, 1.6, WHITE, radius=0.04)
techs = [
    ("Python / FastAPI", "Backend server"),
    ("Chart.js", "Visualizations"),
    ("Paramiko", "SSH connections"),
    ("YAML Storage", "No database needed"),
    ("Any OS", "Windows · Linux · Mac"),
]
for i, (label, sub) in enumerate(techs):
    x = 1.1 + i * 2.3
    add_text(slide, label, x, 5.0, 2.0, 0.4, size=15, color=DEEP, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, sub, x, 5.4, 2.0, 0.35, size=11, color=GRAY, align=PP_ALIGN.CENTER)

# =================== SLIDE 6 — IMPACT ===================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, DEEP)

add_text(slide, "Impact", 0.8, 0.4, 6, 0.9, size=40, color=WHITE, bold=True, font=FONT_T)

impacts = [
    ("Data > Guesswork", 'Turn "Is the customer using it?" from a phone call into a dashboard'),
    ("Reclaim Idle Systems", "Identify hardware sitting unused — reallocate to active evaluations"),
    ("Prove ROI", "Quantify customer engagement hours for business justification"),
    ("Zero Customer Friction", "No software to install, no firewall changes — just existing SSH access"),
]
for i, (stat, desc) in enumerate(impacts):
    y = 1.6 + i * 1.3
    add_rect(slide, 0.8, y, 11.7, 1.05, NAVY, radius=0.03)
    add_text(slide, stat, 1.1, y + 0.1, 4.5, 0.4, size=22, color=CORAL, bold=True, font=FONT_T)
    add_text(slide, desc, 1.1, y + 0.55, 10.5, 0.4, size=15, color=ICE)

add_text(slide, "UsageLens — Built by FAEs, for FAEs", 0.8, 6.5, 11.7, 0.5, size=20, color=GOLD, bold=True, font=FONT_T, align=PP_ALIGN.CENTER)

# =================== SAVE ===================
out = os.path.join(os.path.dirname(__file__), "UsageLens_Pitch.pptx")
prs.save(out)
print(f"Created: {out}")
